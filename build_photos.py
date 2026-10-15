"""Publish small sheet images as static assets, without browser XLSX downloads."""
import base64, hashlib, io, json, re, sys, time, urllib.request
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
from PIL import Image

DEST = Path(sys.argv[1] if len(sys.argv) > 1 else 'site')
URL = 'https://docs.google.com/spreadsheets/d/1ieuovSRmRkGgOjjTDXpnXwwc-5612l0RkPLKkct0isQ/export?format=xlsx'
NAMES = {'수납샘플','개발샘플','공장샘플','구매샘플','통합시트'}

def key(values):
    return '|'.join(re.sub(r'\s+', ' ', str(values[i] if values[i] is not None else '')).strip() for i in [0,1,4,5,9])

def main():
    DEST.mkdir(parents=True,exist_ok=True)
    (DEST/'photos').mkdir(exist_ok=True)
    for attempt in range(5):
        try:
            with urllib.request.urlopen(URL, timeout=180) as response:
                payload=response.read()
            book=openpyxl.load_workbook(io.BytesIO(payload),data_only=True)
            break
        except Exception as error:
            if attempt==4: raise
            delay=10*(2**attempt)
            print(f'Sheet download interrupted ({type(error).__name__}); retrying in {delay}s.', flush=True)
            time.sleep(delay)
    photos={}
    for sheet in book:
        name=next((label for label in NAMES if label in sheet.title), None)
        if name not in NAMES: continue
        counts={}; row_keys={}
        for row_number in range(2, sheet.max_row+1):
            values=[sheet.cell(row_number,i+1).value for i in range(11)]
            if not any(str(v or '').strip() for v in values[:6]) and not str(values[10] or '').strip():
                continue
            identity=key(values)
            counts[identity]=counts.get(identity,0)+1
            row_keys[row_number]=identity+'|occ:'+str(counts[identity])
        for image in sheet._images:
            anchor=getattr(image.anchor,'_from',None)
            if anchor is None or anchor.col!=6: continue
            values=[sheet.cell(anchor.row+1,i+1).value for i in range(11)]
            if not values[4]: continue
            im=Image.open(io.BytesIO(image._data())).convert('RGBA')
            im.thumbnail((420,420),Image.Resampling.LANCZOS)
            rgb=Image.new('RGB',im.size,'white');rgb.paste(im,mask=im.getchannel('A'))
            encoded=io.BytesIO();rgb.save(encoded,'JPEG',quality=72,optimize=True)
            jpeg=encoded.getvalue();filename=hashlib.sha256(jpeg).hexdigest()[:24]+'.jpg'
            (DEST/'photos'/filename).write_bytes(jpeg)
            src='photos/'+filename
            row_key=row_keys.get(anchor.row+1)
            if row_key:
                photos[name+'|'+row_key]=src
                if name=='통합시트': photos['*|'+row_key]=src
    if not photos: raise RuntimeError('No spreadsheet photos extracted; keeping previous deployed site unchanged.')
    result={'updatedAt':datetime.now(timezone.utc).isoformat(),'photos':photos}
    (DEST/'photos.json').write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
    (DEST/'_headers').write_text('/photos.json\n  Cache-Control: no-cache\n/photos/*\n  Cache-Control: public, max-age=31536000, immutable\n',encoding='utf-8')
    print('Published',len(set(photos.values())),'small images for',len(photos),'sample references.')

if __name__=='__main__': main()
